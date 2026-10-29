#!/usr/bin/env python3
"""Compile War Chronicle CardData + EventFlow worksheets into portable JSON.

Usage (if this script is beside War Chronicle Master.xlsx):
    python wc_compile.py

Or specify files explicitly:
    python wc_compile.py --input "War Chronicle Master.xlsx" --output "wc_data.json"

This compiler is intentionally platform-neutral. It contains no Tabletop Simulator
logic. TTS, a browser app, or a future standalone implementation can consume the
same JSON and provide its own adapter/executor.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

SCHEMA_VERSION = "0.3.0"

CARD_SHEET = "CardData"
FLOW_SHEET = "EventFlow"
EXPLORE_SHEET = "explore-cards"
ARRIVAL_SHEET = "arrival-cards"
ADVANCEMENT_SHEET = "advancement-cards"

CARD_HEADERS = [
    "CardID", "CardType", "Set", "CardNum", "EchoPool", "EchoNum",
    "Title", "Keyword", "Location", "Time", "CostType", "CostValue",
    "AfterPlay", "Flavor", "Effect", "Notes",
]


EXPLORE_HEADERS = ["CardNum", "Title", "Terrain", "Image", "Flavor", "Effect"]
ARRIVAL_HEADERS = ["CardNum", "Deck", "Title", "Tribe", "Tribute", "Setup", "Effect", "Host"]
ADVANCEMENT_HEADERS = ["CardID", "Title", "Keyword", "Cost", "Commit", "Effect"]

FLOW_HEADERS = [
    "FlowID", "SourceType", "SourceID", "Trigger", "NodeID", "NodeType",
    "Prompt", "Input", "Branch", "Operator", "Value", "ActionType",
    "ActionTarget", "ActionValue", "NextNode", "Flags", "Notes", "PlayerText",
]

CARD_TYPES = {"Camp", "Scene", "Echo"}
COST_TYPES = {"None", "Fixed", "Variable"}
AFTER_PLAY_VALUES = {"Keep", "Discard", "Remove", "RemoveOrDiscard", "Special"}
NODE_TYPES = {"Choice", "Select", "Roll", "Test", "Check", "Action", "End"}

# Sources that deliberately live outside CardData but can own EventFlow rules.
EXTERNAL_SOURCE_TYPES = {"Explore", "Arrival", "Token", "GameRule", "UnitRule"}

_INT_RE = re.compile(r"^[+-]?\d+$")
_FLOAT_RE = re.compile(r"^[+-]?(?:\d+\.\d*|\d*\.\d+)$")


class CompileError(Exception):
    pass


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def clean_identifier(value: Any) -> str | None:
    if is_blank(value):
        return None
    return str(value).strip()


def coerce_scalar(value: Any) -> Any:
    """Coerce plain numeric strings while preserving semantic strings/ranges."""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            return None
        if _INT_RE.fullmatch(text):
            try:
                return int(text)
            except ValueError:
                pass
        if _FLOAT_RE.fullmatch(text):
            try:
                return float(text)
            except ValueError:
                pass
        return text
    return value


def parse_flags(value: Any) -> list[dict[str, Any]]:
    """Parse semicolon-delimited flags without requiring a rigid flag schema."""
    if is_blank(value):
        return []
    result: list[dict[str, Any]] = []
    for token in str(value).split(";"):
        token = token.strip()
        if not token:
            continue
        if "=" in token:
            name, raw_value = token.split("=", 1)
            result.append({"name": name.strip(), "value": coerce_scalar(raw_value)})
        else:
            result.append({"name": token, "value": True})
    return result


def read_sheet_rows(ws, required_headers: list[str]) -> list[dict[str, Any]]:
    raw_headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = [str(h).strip() if h is not None else None for h in raw_headers]

    missing = [h for h in required_headers if h not in headers]
    if missing:
        raise CompileError(
            f"Worksheet '{ws.title}' is missing required columns: {', '.join(missing)}"
        )

    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(not is_blank(v) for v in values):
            continue
        record = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        record["__row__"] = row_number
        rows.append(record)
    return rows


def normalize_card(row: dict[str, Any]) -> dict[str, Any]:
    cost_type = clean_identifier(row.get("CostType"))
    cost_value = coerce_scalar(row.get("CostValue"))

    return {
        "id": clean_identifier(row.get("CardID")),
        "type": clean_identifier(row.get("CardType")),
        "set": clean_identifier(row.get("Set")),
        "card_num": coerce_scalar(row.get("CardNum")),
        "echo_pool": clean_identifier(row.get("EchoPool")),
        "echo_num": coerce_scalar(row.get("EchoNum")),
        "title": row.get("Title"),
        "keyword": clean_identifier(row.get("Keyword")),
        "location": clean_identifier(row.get("Location")),
        "time": coerce_scalar(row.get("Time")),
        "cost": {
            "type": cost_type,
            "value": cost_value,
        },
        "after_play": clean_identifier(row.get("AfterPlay")),
        "text": {
            "flavor": row.get("Flavor"),
            "effect": row.get("Effect"),
        },
        "notes": row.get("Notes"),
    }


def compile_cards(rows: list[dict[str, Any]], errors: list[str], warnings: list[str]):
    cards: OrderedDict[str, dict[str, Any]] = OrderedDict()
    printed_nums: dict[Any, str] = {}
    echo_keys: dict[tuple[str, Any], str] = {}

    for row in rows:
        rnum = row["__row__"]
        card = normalize_card(row)
        cid = card["id"]
        ctype = card["type"]

        if not cid:
            errors.append(f"{CARD_SHEET} row {rnum}: CardID is required.")
            continue
        if cid in cards:
            errors.append(f"{CARD_SHEET} row {rnum}: duplicate CardID '{cid}'.")
            continue
        if ctype not in CARD_TYPES:
            errors.append(f"{CARD_SHEET} row {rnum} ({cid}): unknown CardType '{ctype}'.")

        if ctype in {"Camp", "Scene"}:
            if card["card_num"] is None:
                errors.append(f"{CARD_SHEET} row {rnum} ({cid}): CardNum is required for {ctype} cards.")
            elif card["card_num"] in printed_nums:
                errors.append(
                    f"{CARD_SHEET} row {rnum} ({cid}): CardNum {card['card_num']} already used by "
                    f"{printed_nums[card['card_num']]} ."
                )
            else:
                printed_nums[card["card_num"]] = cid
            if card["echo_pool"] is not None or card["echo_num"] is not None:
                errors.append(f"{CARD_SHEET} row {rnum} ({cid}): non-Echo card must not define EchoPool/EchoNum.")

        if ctype == "Echo":
            if card["card_num"] is not None:
                errors.append(f"{CARD_SHEET} row {rnum} ({cid}): Echo CardNum must be blank.")
            if not card["echo_pool"] or card["echo_num"] is None:
                errors.append(f"{CARD_SHEET} row {rnum} ({cid}): EchoPool and EchoNum are required for Echo cards.")
            else:
                ekey = (card["echo_pool"], card["echo_num"])
                if ekey in echo_keys:
                    errors.append(
                        f"{CARD_SHEET} row {rnum} ({cid}): duplicate EchoPool/EchoNum {ekey}; "
                        f"already used by {echo_keys[ekey]}."
                    )
                else:
                    echo_keys[ekey] = cid

        cost_type = card["cost"]["type"]
        cost_value = card["cost"]["value"]
        if cost_type is not None and cost_type not in COST_TYPES:
            warnings.append(f"{CARD_SHEET} row {rnum} ({cid}): unfamiliar CostType '{cost_type}'.")
        if cost_type == "Fixed" and cost_value is None:
            errors.append(f"{CARD_SHEET} row {rnum} ({cid}): Fixed cost requires CostValue.")
        if cost_type in {"None", "Variable"} and cost_value is not None:
            warnings.append(
                f"{CARD_SHEET} row {rnum} ({cid}): CostType {cost_type} normally has blank CostValue; "
                f"found {cost_value!r}."
            )

        after_play = card["after_play"]
        if after_play is not None and after_play not in AFTER_PLAY_VALUES:
            warnings.append(f"{CARD_SHEET} row {rnum} ({cid}): unfamiliar AfterPlay '{after_play}'.")

        cards[cid] = card

    echo_pools: OrderedDict[str, list[str]] = OrderedDict()
    pool_members: defaultdict[str, list[tuple[Any, str]]] = defaultdict(list)
    for cid, card in cards.items():
        if card["type"] == "Echo" and card["echo_pool"]:
            pool_members[card["echo_pool"]].append((card["echo_num"], cid))

    for pool in sorted(pool_members):
        def sort_key(item):
            num, cid = item
            return (0, num) if isinstance(num, (int, float)) else (1, str(num), cid)
        echo_pools[pool] = [cid for _, cid in sorted(pool_members[pool], key=sort_key)]

    return cards, echo_pools


def compile_explore_cards(rows: list[dict[str, Any]], errors: list[str]):
    cards: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row in rows:
        rnum = row["__row__"]
        raw_num = coerce_scalar(row.get("CardNum"))
        if raw_num is None:
            errors.append(f"{EXPLORE_SHEET} row {rnum}: CardNum is required.")
            continue
        try:
            num = int(raw_num)
        except (TypeError, ValueError):
            errors.append(f"{EXPLORE_SHEET} row {rnum}: CardNum '{raw_num}' must be numeric.")
            continue
        eid = f"EXPLORE-{num:03d}"
        if eid in cards:
            errors.append(f"{EXPLORE_SHEET} row {rnum}: duplicate Explore card {eid}.")
            continue
        terrain_raw = clean_identifier(row.get("Terrain")) or ""
        terrain_keywords = [t.strip() for t in terrain_raw.split("::") if t.strip()]
        cards[eid] = {
            "id": eid,
            "card_num": num,
            "title": row.get("Title"),
            "terrain_keywords": terrain_keywords,
            "image": clean_identifier(row.get("Image")),
            "text": {
                "flavor": row.get("Flavor"),
                "effect": row.get("Effect"),
            },
        }
    return cards


def compile_arrival_cards(rows: list[dict[str, Any]], errors: list[str]):
    cards: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row in rows:
        rnum = row["__row__"]
        raw_num = coerce_scalar(row.get("CardNum"))
        if raw_num is None:
            errors.append(f"{ARRIVAL_SHEET} row {rnum}: CardNum is required.")
            continue
        try:
            num = int(raw_num)
        except (TypeError, ValueError):
            errors.append(f"{ARRIVAL_SHEET} row {rnum}: CardNum '{raw_num}' must be numeric.")
            continue
        aid = f"ARRIVAL-{num:03d}"
        if aid in cards:
            errors.append(f"{ARRIVAL_SHEET} row {rnum}: duplicate Arrival card {aid}.")
            continue
        cards[aid] = {
            "id": aid,
            "card_num": num,
            "deck": clean_identifier(row.get("Deck")),
            "title": row.get("Title"),
            "tribe": clean_identifier(row.get("Tribe")),
            "tribute": clean_identifier(row.get("Tribute")),
            "setup": row.get("Setup"),
            "effect": row.get("Effect"),
            "host": row.get("Host"),
        }
    return cards


def compile_advancements(rows: list[dict[str, Any]], errors: list[str]):
    cards: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row in rows:
        rnum = row["__row__"]
        raw_id = clean_identifier(row.get("CardID"))
        if not raw_id:
            errors.append(f"{ADVANCEMENT_SHEET} row {rnum}: CardID is required.")
            continue
        if raw_id.isdigit():
            aid = f"ADV-{int(raw_id):03d}"
            card_num: Any = int(raw_id)
        else:
            aid = f"ADV-{raw_id}"
            card_num = raw_id
        if aid in cards:
            errors.append(f"{ADVANCEMENT_SHEET} row {rnum}: duplicate Advancement card '{aid}'.")
            continue
        cards[aid] = {
            "id": aid,
            "card_num": card_num,
            "title": row.get("Title"),
            "keyword": clean_identifier(row.get("Keyword")),
            "cost": clean_identifier(row.get("Cost")),
            "commit": clean_identifier(row.get("Commit")),
            "effect": row.get("Effect"),
        }
    return cards


def collect_flow_groups(rows: list[dict[str, Any]], errors: list[str]):
    groups: OrderedDict[str, list[dict[str, Any]]] = OrderedDict()
    for row in rows:
        fid = clean_identifier(row.get("FlowID"))
        if not fid:
            errors.append(f"{FLOW_SHEET} row {row['__row__']}: FlowID is required.")
            continue
        groups.setdefault(fid, []).append(row)
    return groups


def one_consistent(rows: Iterable[dict[str, Any]], field: str, flow_id: str, errors: list[str]) -> str | None:
    values = {clean_identifier(r.get(field)) for r in rows}
    values.discard(None)
    if len(values) > 1:
        errors.append(f"{FLOW_SHEET} flow {flow_id}: inconsistent {field} values: {sorted(values)}")
    return next(iter(values), None)


def compile_node(node_id: str, rows: list[dict[str, Any]], flow_id: str,
                 errors: list[str], warnings: list[str]) -> dict[str, Any]:
    types = {clean_identifier(r.get("NodeType")) for r in rows}
    types.discard(None)
    if len(types) != 1:
        errors.append(f"{FLOW_SHEET} flow {flow_id}, node {node_id}: expected one NodeType; found {sorted(types)}")
        node_type = next(iter(types), None)
    else:
        node_type = next(iter(types))

    if node_type not in NODE_TYPES:
        warnings.append(f"{FLOW_SHEET} flow {flow_id}, node {node_id}: unfamiliar NodeType '{node_type}'.")

    prompts = {r.get("Prompt") for r in rows if not is_blank(r.get("Prompt"))}
    if len(prompts) > 1:
        warnings.append(f"{FLOW_SHEET} flow {flow_id}, node {node_id}: multiple Prompt values found.")
    prompt = next(iter(prompts), None)

    node: dict[str, Any] = {"type": node_type}
    if prompt is not None:
        node["prompt"] = prompt

    if node_type == "Action":
        if len(rows) != 1:
            errors.append(f"{FLOW_SHEET} flow {flow_id}, Action node {node_id}: expected exactly one row.")
        row = rows[0]
        node["action"] = {
            "type": clean_identifier(row.get("ActionType")),
            "target": clean_identifier(row.get("ActionTarget")),
            "value": coerce_scalar(row.get("ActionValue")),
        }
        if not node["action"]["type"]:
            errors.append(f"{FLOW_SHEET} flow {flow_id}, Action node {node_id}: ActionType is required.")
        nxt = clean_identifier(row.get("NextNode"))
        if nxt:
            node["next"] = nxt
        flags = parse_flags(row.get("Flags"))
        if flags:
            node["flags"] = flags
        if not is_blank(row.get("Notes")):
            node["notes"] = row.get("Notes")
        if not is_blank(row.get("PlayerText")):
            node["player_text"] = row.get("PlayerText")
        return node

    if node_type == "End":
        if len(rows) != 1:
            warnings.append(f"{FLOW_SHEET} flow {flow_id}, End node {node_id}: expected one row; found {len(rows)}.")
        return node

    # Branching nodes: Choice, Select, Roll, Test, Check.
    branches = []
    for row in rows:
        branch = {
            "label": row.get("Branch"),
            "input": clean_identifier(row.get("Input")),
            "operator": clean_identifier(row.get("Operator")),
            "value": coerce_scalar(row.get("Value")),
            "next": clean_identifier(row.get("NextNode")),
        }
        flags = parse_flags(row.get("Flags"))
        if flags:
            branch["flags"] = flags
        if not is_blank(row.get("Notes")):
            branch["notes"] = row.get("Notes")
        if not is_blank(row.get("PlayerText")):
            branch["player_text"] = row.get("PlayerText")
        branches.append(branch)
    node["branches"] = branches
    return node


def compile_flows(rows: list[dict[str, Any]], cards: OrderedDict[str, dict[str, Any]],
                  echo_pools: OrderedDict[str, list[str]], errors: list[str], warnings: list[str]):
    groups = collect_flow_groups(rows, errors)
    flows: OrderedDict[str, dict[str, Any]] = OrderedDict()
    source_index: OrderedDict[str, dict[str, Any]] = OrderedDict()
    external_sources: OrderedDict[tuple[str, str], dict[str, str]] = OrderedDict()

    for flow_id, flow_rows in groups.items():
        source_type = one_consistent(flow_rows, "SourceType", flow_id, errors)
        source_id = one_consistent(flow_rows, "SourceID", flow_id, errors)
        trigger = one_consistent(flow_rows, "Trigger", flow_id, errors)

        if not source_type or not source_id or not trigger:
            errors.append(f"{FLOW_SHEET} flow {flow_id}: SourceType, SourceID, and Trigger are required.")

        if source_id in cards:
            card_type = cards[source_id]["type"]
            if source_type != card_type:
                warnings.append(
                    f"{FLOW_SHEET} flow {flow_id}: SourceType '{source_type}' does not match "
                    f"CardData type '{card_type}' for {source_id}."
                )
        elif source_type in EXTERNAL_SOURCE_TYPES and source_id:
            external_sources[(source_type, source_id)] = {"type": source_type, "id": source_id}
        else:
            errors.append(
                f"{FLOW_SHEET} flow {flow_id}: SourceID '{source_id}' is not in CardData and "
                f"SourceType '{source_type}' is not an allowed external source type."
            )

        node_rows: OrderedDict[str, list[dict[str, Any]]] = OrderedDict()
        for row in flow_rows:
            nid = clean_identifier(row.get("NodeID"))
            if not nid:
                errors.append(f"{FLOW_SHEET} row {row['__row__']} flow {flow_id}: NodeID is required.")
                continue
            node_rows.setdefault(nid, []).append(row)

        nodes: OrderedDict[str, dict[str, Any]] = OrderedDict()
        for node_id, nrows in node_rows.items():
            nodes[node_id] = compile_node(node_id, nrows, flow_id, errors, warnings)

        node_ids = set(nodes)
        referenced = set()
        for node in nodes.values():
            if node.get("next"):
                referenced.add(node["next"])
            for branch in node.get("branches", []):
                if branch.get("next"):
                    referenced.add(branch["next"])

        missing_targets = sorted(referenced - node_ids)
        if missing_targets:
            errors.append(
                f"{FLOW_SHEET} flow {flow_id}: NextNode references missing node(s): {', '.join(missing_targets)}"
            )

        # N01 is the conventional entry point. If absent, infer the sole root node.
        if "N01" in node_ids:
            start_node = "N01"
        else:
            roots = [nid for nid in nodes if nid not in referenced]
            if len(roots) == 1:
                start_node = roots[0]
            elif len(roots) == 0:
                errors.append(
                    f"{FLOW_SHEET} flow {flow_id}: no N01 and no inferable root node (all nodes are referenced)."
                )
                start_node = next(iter(nodes), None)
            else:
                errors.append(
                    f"{FLOW_SHEET} flow {flow_id}: no N01 and multiple possible root nodes: {', '.join(roots)}"
                )
                start_node = roots[0]

        # Reachability check; loops are fine.
        if start_node and start_node in nodes:
            reachable = set()
            stack = [start_node]
            while stack:
                nid = stack.pop()
                if nid in reachable or nid not in nodes:
                    continue
                reachable.add(nid)
                node = nodes[nid]
                if node.get("next"):
                    stack.append(node["next"])
                for branch in node.get("branches", []):
                    if branch.get("next"):
                        stack.append(branch["next"])
            unreachable = [nid for nid in nodes if nid not in reachable]
            if unreachable:
                warnings.append(
                    f"{FLOW_SHEET} flow {flow_id}: unreachable node(s) from {start_node}: {', '.join(unreachable)}"
                )

        # Semantic reference checks for Echo actions.
        for node_id, node in nodes.items():
            if node.get("type") != "Action":
                continue
            action = node.get("action", {})
            atype = action.get("type")
            target = action.get("target")
            value = action.get("value")
            if atype in {"SeedRandomEcho", "SeedRemainingEcho"} and target == "EchoPool":
                if str(value) not in echo_pools:
                    errors.append(
                        f"{FLOW_SHEET} flow {flow_id}, node {node_id}: EchoPool '{value}' does not exist in CardData."
                    )
            if atype == "SeedEcho" and target == "CardID":
                if value not in cards:
                    errors.append(
                        f"{FLOW_SHEET} flow {flow_id}, node {node_id}: referenced CardID '{value}' does not exist."
                    )
                elif cards[value]["type"] != "Echo":
                    errors.append(
                        f"{FLOW_SHEET} flow {flow_id}, node {node_id}: SeedEcho target '{value}' is not an Echo card."
                    )

        flows[flow_id] = {
            "source": {"type": source_type, "id": source_id},
            "trigger": trigger,
            "start_node": start_node,
            "nodes": nodes,
        }

        if source_id:
            entry = source_index.setdefault(source_id, {"source_type": source_type, "triggers": OrderedDict()})
            entry["triggers"].setdefault(trigger, []).append(flow_id)

    return flows, source_index, list(external_sources.values())


def discovered_vocabulary(flows: OrderedDict[str, dict[str, Any]]) -> dict[str, list[str]]:
    node_types = set()
    action_types = set()
    triggers = set()
    source_types = set()
    operators = set()

    for flow in flows.values():
        triggers.add(flow["trigger"])
        source_types.add(flow["source"]["type"])
        for node in flow["nodes"].values():
            if node.get("type"):
                node_types.add(node["type"])
            action = node.get("action")
            if action and action.get("type"):
                action_types.add(action["type"])
            for branch in node.get("branches", []):
                if branch.get("operator"):
                    operators.add(branch["operator"])

    return {
        "source_types": sorted(source_types),
        "triggers": sorted(triggers),
        "node_types": sorted(node_types),
        "action_types": sorted(action_types),
        "operators": sorted(operators),
    }


def compile_workbook(input_path: Path) -> tuple[dict[str, Any], list[str], list[str]]:
    if not input_path.exists():
        raise CompileError(f"Workbook not found: {input_path}")

    wb = load_workbook(input_path, read_only=True, data_only=False)
    try:
        for sheet in (CARD_SHEET, FLOW_SHEET, EXPLORE_SHEET, ARRIVAL_SHEET, ADVANCEMENT_SHEET):
            if sheet not in wb.sheetnames:
                raise CompileError(f"Workbook is missing required worksheet '{sheet}'.")

        errors: list[str] = []
        warnings: list[str] = []

        card_rows = read_sheet_rows(wb[CARD_SHEET], CARD_HEADERS)
        flow_rows = read_sheet_rows(wb[FLOW_SHEET], FLOW_HEADERS)
        explore_rows = read_sheet_rows(wb[EXPLORE_SHEET], EXPLORE_HEADERS)
        arrival_rows = read_sheet_rows(wb[ARRIVAL_SHEET], ARRIVAL_HEADERS)
        advancement_rows = read_sheet_rows(wb[ADVANCEMENT_SHEET], ADVANCEMENT_HEADERS)

        cards, echo_pools = compile_cards(card_rows, errors, warnings)
        explore_cards = compile_explore_cards(explore_rows, errors)
        arrival_cards = compile_arrival_cards(arrival_rows, errors)
        advancements = compile_advancements(advancement_rows, errors)
        flows, source_index, external_sources = compile_flows(
            flow_rows, cards, echo_pools, errors, warnings
        )

        payload = {
            "schema_version": SCHEMA_VERSION,
            "generated_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "source_workbook": input_path.name,
            "cards": cards,
            "explore_cards": explore_cards,
            "arrival_cards": arrival_cards,
            "advancements": advancements,
            "echo_pools": echo_pools,
            "flows": flows,
            "source_index": source_index,
            "external_sources": external_sources,
            "vocabulary": discovered_vocabulary(flows),
        }

        return payload, errors, warnings
    finally:
        wb.close()


def parse_args() -> argparse.Namespace:
    default_input = Path(__file__).resolve().with_name("War Chronicle Master.xlsx")
    parser = argparse.ArgumentParser(description="Compile War Chronicle Excel data into portable JSON.")
    parser.add_argument(
        "--input", "-i", type=Path, default=default_input,
        help=f"Input .xlsx workbook (default: {default_input.name} beside this script)",
    )
    parser.add_argument(
        "--output", "-o", type=Path, default=None,
        help="Output JSON path (default: wc_data.json beside the input workbook)",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    output_path = (
        args.output.expanduser().resolve()
        if args.output is not None
        else input_path.with_name("wc_data.json")
    )

    try:
        payload, errors, warnings = compile_workbook(input_path)
    except CompileError as exc:
        print(f"ERROR: {exc}")
        return 1
    except Exception as exc:
        print(f"UNEXPECTED ERROR: {exc}")
        return 1

    print(f"Workbook: {input_path}")
    print(f"Cards: {len(payload['cards'])}")
    print(f"Explore cards: {len(payload['explore_cards'])}")
    print(f"Arrival cards: {len(payload['arrival_cards'])}")
    print(f"Advancements: {len(payload['advancements'])}")
    print(f"Echo pools: {len(payload['echo_pools'])}")
    print(f"Flows: {len(payload['flows'])}")
    print(f"External flow sources: {len(payload['external_sources'])}")

    if warnings:
        print(f"\nWarnings ({len(warnings)}):")
        for warning in warnings:
            print(f"  - {warning}")

    if errors:
        print(f"\nERRORS ({len(errors)}): JSON was NOT written.")
        for error in errors:
            print(f"  - {error}")
        return 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    print("\nValidation passed.")
    print(f"Created: {output_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
